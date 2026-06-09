"""Postprocess thesis docx for MoldSU formatting requirements.

Fixes applied (in order):
1. Page size: A4 (21 × 29.7 cm)  (R3_4 guide §3.1)
2. Font size: Normal body text 14pt → 12pt  (R3_4 guide §3.2)
3. TOC heading: 'Table of Contents' → 'СОДЕРЖАНИЕ'
4. Titulnyi pages: generate RO + RU title pages from metadata.yaml
5. Margins: left 2.5 cm, right 1.5 cm, top 2 cm, bottom 2 cm
6. Page numbers: bottom center
"""
from __future__ import annotations

import argparse
import copy
from pathlib import Path

import yaml
import docx
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH


# МолдГУ page size and margins
PAGE_WIDTH = Cm(21.0)
PAGE_HEIGHT = Cm(29.7)
MARGIN_LEFT = Cm(2.5)
MARGIN_RIGHT = Cm(1.5)
MARGIN_TOP = Cm(2.0)
MARGIN_BOTTOM = Cm(2.0)


# ── Page size ─────────────────────────────────────────────────────────────────

def set_a4_page_size(document: Document) -> None:
    """Set all sections to A4 (21 × 29.7 cm) with МолдГУ margins."""
    for section in document.sections:
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
    print("  page size → A4 (21×29.7 cm)")


# ── TOC heading fix ───────────────────────────────────────────────────────────

def fix_toc_heading(document: Document) -> None:
    """Replace 'Table of Contents' with 'СОДЕРЖАНИЕ' in all runs."""
    replaced = False
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if "Table of Contents" in run.text:
                run.text = run.text.replace("Table of Contents", "СОДЕРЖАНИЕ")
                replaced = True
    if replaced:
        print("  TOC heading → 'СОДЕРЖАНИЕ'")


# ── Font size fix ─────────────────────────────────────────────────────────────

def fix_font_size(document: Document, target_pt: int = 12) -> None:
    """Set all non-Heading-1 text runs to target_pt.

    Heading 1 (chapter titles) must stay at 14 pt per MoldSU §3.2.
    All other paragraphs (Normal, Heading 2/3, captions, table cells)
    are set to target_pt so they override whatever the reference template
    carried in each run's direct formatting.
    """
    target_size = Pt(target_pt)

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else "Normal"
        if style_name == "Heading 1":
            continue
        for run in paragraph.runs:
            run.font.size = target_size

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = target_size

    # Update the Normal style default so any unstyled runs also inherit 12pt.
    for style in document.styles:
        if style.name == "Normal":
            style.font.size = target_size
            break

    print(f"  font size → {target_pt}pt on all non-Heading-1 runs")


# ── Title page helpers ────────────────────────────────────────────────────────

def _add_para(doc: Document, text: str, *,
               bold: bool = False,
               pt: int = 12,
               align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.CENTER,
               space_before: int = 0,
               space_after: int = 0) -> docx.text.paragraph.Paragraph:
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(pt)
    return para


def _add_page_break(doc: Document) -> None:
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE)


def _build_title_page_ro(doc: Document, meta: dict) -> None:
    """Romanian title page (Appendix 4 MoldSU template)."""
    _add_para(doc, "UNIVERSITATEA DE STAT DIN MOLDOVA", bold=True, pt=12, space_after=0)
    _add_para(doc, "FACULTATEA DE MATEMATICĂ ȘI INFORMATICĂ", bold=True, pt=12)
    _add_para(doc, "DEPARTAMENTUL DE INFORMATICĂ", bold=True, pt=12, space_after=24)

    _add_para(doc, "")  # spacer

    author = meta.get("author_ro", "PRENUME NUME")
    _add_para(doc, author, bold=True, pt=16, space_before=12, space_after=12)

    _add_para(doc, "")  # spacer

    title = meta.get("title_ro", "TITLUL TEZEI")
    _add_para(doc, title, bold=True, pt=14, space_before=6, space_after=6)

    domain = meta.get("domain_ro", meta.get("domain", "0613. Dezvoltarea de software și aplicații"))
    _add_para(doc, f"Domeniul de formare profesională: {domain}", pt=12)

    program = meta.get("program_ro", "Informatică Aplicată")
    _add_para(doc, f"Programul: {program}", pt=12, space_after=24)

    _add_para(doc, "")  # spacer

    _add_para(doc, "Teză de master", bold=True, pt=14, space_after=48)

    # Signature block — left-aligned
    dept_head = meta.get("dept_head_name", "Prenume Nume")
    dept_head_title = meta.get("dept_head_title_ro", "dr., conf. univ.")
    supervisor = meta.get("supervisor_name", "Prenume Nume")
    supervisor_title = meta.get("supervisor_title_ro", "dr., prof. univ.")
    author_ro = meta.get("author_ro", "Prenume Nume")

    _add_para(doc, f"Șef de departament:   ________   {dept_head}, {dept_head_title}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(semnătura)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=12)

    _add_para(doc, f"Conducător științific:   ________   {supervisor}, {supervisor_title}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(semnătura)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=12)

    _add_para(doc, f"Autor:   ________   {author_ro}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(semnătura)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=36)

    year = meta.get("year", 2026)
    _add_para(doc, f"Chișinău — {year}", bold=True, pt=12)


def _build_title_page_ru(doc: Document, meta: dict) -> None:
    """Russian title page (second titulnyi, same layout in Russian)."""
    _add_para(doc, "МОЛДАВСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ", bold=True, pt=12, space_after=0)
    _add_para(doc, "ФАКУЛЬТЕТ МАТЕМАТИКИ И ИНФОРМАТИКИ", bold=True, pt=12)
    _add_para(doc, "ДЕПАРТАМЕНТ ИНФОРМАТИКИ", bold=True, pt=12, space_after=24)

    _add_para(doc, "")

    author = meta.get("author_ru", "ИМЯ ФАМИЛИЯ")
    _add_para(doc, author, bold=True, pt=16, space_before=12, space_after=12)

    _add_para(doc, "")

    title = meta.get("title_ru", "НАЗВАНИЕ ДИССЕРТАЦИИ")
    _add_para(doc, title, bold=True, pt=14, space_before=6, space_after=6)

    domain = meta.get("domain", "0613. Разработка программного обеспечения и приложений")
    _add_para(doc, f"Область профессиональной подготовки: {domain}", pt=12)

    program = meta.get("program_ru", "Прикладная информатика")
    _add_para(doc, f"Программа: {program}", pt=12, space_after=24)

    _add_para(doc, "")

    _add_para(doc, "Магистерская работа", bold=True, pt=14, space_after=48)

    dept_head = meta.get("dept_head_name", "Имя Фамилия")
    dept_head_title = meta.get("dept_head_title_ru", "доктор, доцент")
    supervisor = meta.get("supervisor_name", "Имя Фамилия")
    supervisor_title = meta.get("supervisor_title_ru", "доктор, профессор")
    author_ru = meta.get("author_ru", "Имя Фамилия")

    _add_para(doc, f"Зав. кафедрой:   ________   {dept_head}, {dept_head_title}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(подпись)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=12)

    _add_para(doc, f"Научный руководитель:   ________   {supervisor}, {supervisor_title}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(подпись)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=12)

    _add_para(doc, f"Автор:   ________   {author_ru}",
              align=WD_ALIGN_PARAGRAPH.LEFT, pt=12)
    _add_para(doc, "(подпись)", align=WD_ALIGN_PARAGRAPH.LEFT, pt=10, space_after=36)

    year = meta.get("year", 2026)
    _add_para(doc, f"Кишинёв — {year}", bold=True, pt=12)


# ── Document merge ────────────────────────────────────────────────────────────

def _copy_element(element):
    """Deep-copy an XML element so it can be inserted into another document."""
    return copy.deepcopy(element)


def prepend_titulnyi(main_doc: Document, meta: dict) -> Document:
    """Prepend 2 title pages to main_doc by inserting XML at the start of body.

    Mutates main_doc in-place so that all embedded parts (images, styles,
    relationships) are preserved.  A new Document() would lose those parts.

    Uses addprevious() (lxml sibling insertion) to place title page elements
    before the first existing body child.
    """
    # Build title pages in a scratch Document to get their XML elements.
    tmp = Document()
    tmp_body = tmp.element.body
    for child in list(tmp_body):
        tmp_body.remove(child)

    _build_title_page_ro(tmp, meta)
    _add_page_break(tmp)
    _build_title_page_ru(tmp, meta)
    _add_page_break(tmp)

    title_elements = [_copy_element(child) for child in list(tmp.element.body)]

    main_body = main_doc.element.body
    first_child = main_body[0] if len(main_body) > 0 else None

    if first_child is not None:
        # addprevious() always inserts before first_child (which stays constant),
        # so forward iteration gives the correct final order in the body.
        for elem in title_elements:
            first_child.addprevious(elem)
    else:
        for elem in title_elements:
            main_body.append(elem)

    print("  titulnyi → 2 pages prepended (RO + RU)")
    return main_doc


# ── Margins & page numbers ────────────────────────────────────────────────────

def apply_moldsu_margins(document: Document) -> None:
    for section in document.sections:
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
    print("  margins → left=2.5cm  right=1.5cm  top=2.0cm  bottom=2.0cm")


def _make_page_number_field() -> OxmlElement:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run = OxmlElement("w:r")
    run.append(fld_begin)
    run.append(instr)
    run.append(fld_end)
    return run


def ensure_page_numbers(document: Document) -> None:
    for section in document.sections:
        footer = section.footer
        has_content = any(p.text.strip() for p in footer.paragraphs)
        if has_content:
            print("  page numbers: already present in footer")
            return
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.clear()
        para._element.append(_make_page_number_field())
    print("  page numbers: added to footer (centered)")


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(description="Postprocess thesis.docx for МолдГУ")
    parser.add_argument("--input", required=True, help="Raw pandoc output docx")
    parser.add_argument("--output", required=True, help="Final thesis.docx path")
    parser.add_argument("--metadata", required=True, help="metadata.yaml path")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    meta_path = Path(args.metadata)

    if not input_path.exists():
        raise SystemExit(f"Input not found: {input_path}")
    if not meta_path.exists():
        raise SystemExit(f"Metadata not found: {meta_path}")

    with meta_path.open(encoding="utf-8") as f:
        meta = yaml.safe_load(f)

    document = docx.Document(str(input_path))

    print("  Applying МолдГУ formatting patches:")

    set_a4_page_size(document)
    fix_font_size(document, target_pt=12)
    fix_toc_heading(document)
    document = prepend_titulnyi(document, meta)
    apply_moldsu_margins(document)
    ensure_page_numbers(document)

    document.save(str(output_path))
    size_kb = output_path.stat().st_size // 1024
    print(f"  Saved: {output_path}  ({size_kb} KB)")


if __name__ == "__main__":
    main()
